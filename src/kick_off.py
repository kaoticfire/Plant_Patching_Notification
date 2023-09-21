""" Email when about to start with plant patching. """
from datetime import datetime as dt
from dateutil.relativedelta import relativedelta
from config import Config as cfg
from filelog import file_log
from mail import __mail_notification
from openpyxl import load_workbook

workbook = ''
records = []
cycle = f'Week {int(cfg.WEEK_NUMBER) + 1}'
try:
    workbook = load_workbook(cfg.FILE)
    worksheet = workbook.active
    for row in worksheet.iter_rows(2):
        for cell in row:
            if cell.value == cycle:
                records.append(worksheet.cell(row=cell.row, column=1).value)

except NameError:
    pass

if dt.today() <= dt(int(dt.today().year), int(dt.today().month), cfg.DAYS):
    month = dt.date(dt.today() - relativedelta(months=1))
else:
    month = dt.today()

if cfg.DEBUG:
    to_who = 'virgil.hoover@cavalry.solutions'
else:
    to_who = 'HistorianSupport@calpine.com; ' \
             'brad.gibson@calpine.com; ' \
             'David.Symons@calpine.com; ' \
             'Laura.Morrical@calpine.com; ' \
             'overwatch@cavalry.solutions'

if cfg.WEEK_NUMBER:
    msg_subj = f'{cfg.CHANGE_TICKET} - Plant Prod Monthly Patching - ' \
               f'{month.strftime("%b %Y")} - Week {cfg.WEEK_NUMBER}'
else:
    msg_subj = f'{cfg.CHANGE_TICKET} - Plant Prod Monthly Patching - ' \
               f'{month.strftime("%b %Y")}'

msg_body= f'<html><body><p>Team,<br />Patching for the ' \
          f'plants will be starting at 16:00 EST for the following sites. ' \
          f'{cfg.SPOKE_TO} at the Help Desk has been notified.<br /><br />' \
          f'{"<br />".join(map(str, sorted(records)))}<br /><br />' \
          f'Virgil Hoover<br />SUPPORT@CAVALRY.SOLUTIONS<br />' \
          f'NETWORK OPERATIONS CENTER ///<br />' \
          f'<a href="tel:+1 (720) 279-2233">+1 (720) 279-2233' \
          f'</a> - 24X7 OVERWATCH</p></body></html>'

__mail_notification(to_who, msg_subj, msg_body)

file_log(f'Change {cfg.CHANGE_TICKET} | Week {cfg.WEEK_NUMBER} | Month '
         f'{month.strftime("%b %Y")} | Spoke With {cfg.SPOKE_TO}')