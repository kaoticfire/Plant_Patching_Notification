""" Email when all clear with plant patching. """
from openpyxl import load_workbook
from datetime import datetime as dt
from dateutil.relativedelta import relativedelta
from time import sleep
from config import Config as cfg
from mail import __mail_notification
from filelog import file_log

workbook = ''
records = []
cycle = f'Week {int(cfg.WEEK_NUMBER) + 1}'
try:
    workbook = load_workbook(cfg.FILE)
    worksheet = workbook.active
    for row in worksheet.iter_rows(2):
        for cell in row:
            if cell.value == cycle:
                records.append(worksheet.cell(row=cell.row, column=1).value)
except FileNotFoundError:
    print('No file found to process.')
    sleep(5)
    quit()

except PermissionError:
    print('This file is locked by another process.')
    sleep(5)
    quit()

except KeyError:
    print('The fle has been corrupted')
    sleep(5)
    quit()

except IOError:
    print('The was an error trying to access the file or a portion of it.')
    sleep(5)
    quit()

except NameError:
    pass

finally:
    workbook.close()

excluded = []
for record in records:
    if excluded:
        if record in excluded:
            records.pop(record)
            
if dt.today() <= dt(int(dt.today().year), int(dt.today().month), cfg.DAYS):
    month = dt.date(dt.today() - relativedelta(months=1))
else:
    month = dt.today()

if cfg.DEBUG:
    to = 'me@example.com'
else:
    to = 'support@example.com'

if cfg.WEEK_NUMBER:
    subj = f'{cfg.CHANGE_TICKET} - Plant Prod Monthly Patching - ' \
                      f'{month.strftime("%b %Y")} - Week {cfg.WEEK_NUMBER}'
else:
    subj = f'{cfg.CHANGE_TICKET} - Plant Prod Monthly Patching - ' \
                     f'{month.strftime("%b %Y")}'
content= f'<html><body><p>Overwatch,<br />All patching for the ' \
         f'plant sites has been completed. Please resume normal' \
         f' monitoring for all involved sites. {cfg.SPOKE_TO} ' \
         f'at the Help Desk has been notified.<br /><br />' \
         f'{"<br />".join(map(str, sorted(records)))}<br /><br />' \
         f'SUPPORT@EXAMPLE.COM</p></body></html>'
__mail_notification(to, subj, content)

for item in records:
    file_log(f'Completion email sent for the {item}. | '
             f'Change {cfg.CHANGE_TICKET} | Week {cfg.WEEK_NUMBER} | '
             f'Spoke With {cfg.SPOKE_TO}')
