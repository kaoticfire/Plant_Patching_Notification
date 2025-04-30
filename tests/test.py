""" Testing functionality of scripts."""
from openpyxl import load_workbook
from os import remove
from win32com.client import Dispatch

FILTER = "[Subject] = 'test'"


def test_excel_communication():
    """ Test I/O along with Excel functionality. """
    _file = 'site.xlsx'
    df = 'Plant'

    wb = load_workbook(_file)
    ws = wb['Sheet']
    value = ws.cell(row=1, column=1)
    assert value.value == df
    wb.close()
    remove(_file)



def test_email_sent():
    """ Test email connection and functionality. """
    outlook = Dispatch('outlook.application')
    message = outlook.CreateItem(0)
    message.to = 'someone@example.com'
    message.Subject = 'test'
    message.body = 'this is a test'
    message.Send()

    mapi = outlook.GetNamespace("MAPI")
    inbox = mapi.GetDefaultFolder(6)
    email = inbox.Items
    for i in email.Restrict(FILTER):
        assert i.body.strip() == 'this is a test'



if __name__ == '__main__':
    test_excel_communication()
    test_email_sent()
